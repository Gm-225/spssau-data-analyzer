from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
import pypdfium2 as pdfium


ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "services" / "questionnaire_pipeline.py"
PLAN = ROOT / "examples" / "generic-questionnaire-plan.json"
PROMPT = "请做完整分析：SAT和VAL为自变量，LOY为因变量；SAT=满意度，VAL=感知价值，LOY=忠诚度。"


def make_fixture(path: Path, seed: int) -> None:
    rng = np.random.default_rng(seed)
    count = 180
    satisfaction = rng.normal(size=count)
    value = 0.25 * satisfaction + rng.normal(scale=0.9, size=count)
    loyalty = 0.60 * satisfaction + 0.35 * value + rng.normal(scale=0.65, size=count)

    def items(latent: np.ndarray, prefix: str) -> dict[str, np.ndarray]:
        return {
            f"{prefix}{index}": np.clip(np.rint(3 + latent + rng.normal(scale=0.55, size=count)), 1, 5).astype(int)
            for index in range(1, 4)
        }

    frame = pd.DataFrame({
        "编号": np.arange(1, count + 1),
        "性别": np.where(np.arange(count) % 2 == 0, "女性", "男性"),
        **items(satisfaction, "SAT"),
        **items(value, "VAL"),
        **items(loyalty, "LOY"),
    })
    frame.to_excel(path, index=False)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--evidence-root", default=str(ROOT / "data" / "validation" / "T-20260720-03-v2-template"))
    parser.add_argument("--seed", type=int, default=20260720)
    args = parser.parse_args()
    evidence_root = Path(args.evidence_root).resolve()
    evidence_root.mkdir(parents=True, exist_ok=True)
    source = evidence_root / "通用问卷合成样例.xlsx"
    output = evidence_root / "report"
    output.mkdir(parents=True, exist_ok=True)
    make_fixture(source, args.seed)

    command = [
        sys.executable,
        str(PIPELINE),
        "--input",
        str(source),
        "--plan",
        str(PLAN),
        "--output-dir",
        str(output),
        "--prompt",
        PROMPT,
        "--seed",
        str(args.seed),
    ]
    completed = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=180)
    if completed.returncode:
        raise AssertionError(completed.stderr or completed.stdout)

    result_path = output / "analysis-result.json"
    docx_path = output / "问卷数据分析报告.docx"
    pdf_path = output / "问卷数据分析报告.pdf"
    for artifact in (result_path, docx_path, pdf_path):
        if not artifact.is_file() or artifact.stat().st_size < 1000:
            raise AssertionError(f"V2 通用报告产物缺失或过小：{artifact}")

    result = json.loads(result_path.read_text(encoding="utf-8"))
    titles = [payload["title"] for payload in result["tables"]]
    required_titles = {"数据基本信息", "缺失值分析", "变量描述及变量分布", "信度检验结果", "维度相关矩阵", "模型摘要", "回归系数", "回归结果说明"}
    missing_titles = sorted(required_titles - set(titles))
    if missing_titles:
        raise AssertionError(f"V2 通用报告缺少模块：{missing_titles}")
    regression = result["raw"].get("regression")
    if not regression or regression["predictors"] != ["SAT", "VAL"] or regression["outcome"] != "LOY":
        raise AssertionError(f"通用研究模型未按计划执行：{regression}")
    if regression["rSquared"] < 0.25:
        raise AssertionError(f"合成基准的回归解释率异常：{regression['rSquared']}")

    document = Document(docx_path)
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "满意度" not in document_text or "忠诚度" not in document_text or "达到 0.05 显著性水平" not in document_text:
        raise AssertionError("Word 报告未包含按实际模型生成的动态论文文字")
    pdf = pdfium.PdfDocument(str(pdf_path))
    if len(pdf) < 2:
        raise AssertionError("PDF 报告页数异常")

    summary = {
        "status": "passed",
        "syntheticOnly": True,
        "sampleRows": 180,
        "dimensions": ["SAT", "VAL", "LOY"],
        "model": "SAT + VAL -> LOY",
        "rSquared": regression["rSquared"],
        "artifacts": {path.name: {"path": str(path), "bytes": path.stat().st_size} for path in (result_path, docx_path, pdf_path)},
        "tableTitles": titles,
    }
    (evidence_root / "machine-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
