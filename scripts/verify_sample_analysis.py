from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import time
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

os.environ["OPENBLAS_NUM_THREADS"] = "1"
os.environ["OMP_NUM_THREADS"] = "1"
os.environ["MKL_NUM_THREADS"] = "1"

import numpy as np
import pandas as pd
import pypdfium2 as pdfium
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SAMPLE = ROOT / "examples" / "synthetic-questionnaire-seed20260717.xlsx"
DEFAULT_CSV = ROOT / "examples" / "synthetic-questionnaire-seed20260717.csv"
PLAN = ROOT / "examples" / "questionnaire-plan.json"
PIPELINE = ROOT / "services" / "questionnaire_pipeline.py"
OUTPUTS = ROOT / "outputs"
FULL_PROMPT = "请对这份问卷数据做完整分析：PP为自变量、AA为因变量，PC和PI为并行中介，PV调节PP→AA直接路径；包括频数、描述统计、信度、效度、相关、Bootstrap中介和调节分析，并按分析报告模板输出。"
LIMITED_PROMPT = "只做问卷信度分析"
INVALID_PROMPT = "请根据这份数据写一首宣传诗"
SYNTHETIC_BENCHMARK = {
    "kmo": 0.8971807579830153,
    "alphas": {
        "PP": 0.9290370042867071,
        "AA": 0.893457899907147,
        "PC": 0.9223309487326403,
        "PI": 0.9454808149932017,
        "PV": 0.9288918429975316,
    },
    "outcomeModel": {
        "const": 1.1033939255207585,
        "PP_c": 0.20422941187812949,
        "PC": 0.20895105718546697,
        "PI": 0.41725991738856594,
        "PV_c": -0.15910984960128077,
        "interaction": 0.12192767776172439,
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def hash_tree(root: Path) -> list[dict[str, Any]]:
    if not root.exists():
        return []
    return [
        {
            "path": str(path.relative_to(ROOT)).replace("\\", "/"),
            "length": path.stat().st_size,
            "sha256": sha256(path),
        }
        for path in sorted(root.rglob("*"))
        if path.is_file()
    ]


def assert_synthetic(path: Path) -> Path:
    resolved = path.resolve()
    examples = (ROOT / "examples").resolve()
    if resolved.parent != examples or "synthetic" not in resolved.name.lower():
        raise SystemExit(f"拒绝非 synthetic 样本：{resolved}")
    if not resolved.exists():
        raise SystemExit(f"synthetic 样本不存在：{resolved}")
    return resolved


def cronbach_alpha(values: np.ndarray) -> float:
    item_variances = values.var(axis=0, ddof=1)
    total_variance = values.sum(axis=1).var(ddof=1)
    count = values.shape[1]
    return float(count / (count - 1) * (1 - item_variances.sum() / total_variance))


def independent_reference(frame: pd.DataFrame) -> dict[str, Any]:
    prefixes = {"PP": 4, "AA": 3, "PC": 4, "PI": 5, "PV": 4}
    items = {
        prefix: [
            column for column in frame.columns
            if str(column).startswith(prefix) and str(column)[len(prefix):len(prefix) + 1].isdigit()
        ]
        for prefix in prefixes
    }
    for prefix, expected in prefixes.items():
        if len(items[prefix]) != expected:
            raise AssertionError(f"独立复核识别 {prefix} 题项数异常：{len(items[prefix])} != {expected}")
    scores = pd.DataFrame({
        prefix: frame[columns].astype(float).mean(axis=1)
        for prefix, columns in items.items()
    })
    all_columns = [column for columns in items.values() for column in columns]
    correlation = frame[all_columns].astype(float).corr().to_numpy()
    inverse = np.linalg.pinv(correlation)
    partial = -inverse / np.sqrt(np.outer(np.diag(inverse), np.diag(inverse)))
    np.fill_diagonal(partial, 0.0)
    correlations_off_diagonal = correlation.copy()
    np.fill_diagonal(correlations_off_diagonal, 0.0)
    kmo = np.square(correlations_off_diagonal).sum() / (
        np.square(correlations_off_diagonal).sum() + np.square(partial).sum()
    )

    centered_pp = scores["PP"] - scores["PP"].mean()
    centered_pv = scores["PV"] - scores["PV"].mean()
    matrix = np.column_stack([
        np.ones(len(scores)),
        centered_pp,
        scores["PC"],
        scores["PI"],
        centered_pv,
        centered_pp * centered_pv,
    ])
    beta = np.linalg.lstsq(matrix, scores["AA"].to_numpy(), rcond=None)[0]
    return {
        "alphas": {
            prefix: cronbach_alpha(frame[columns].astype(float).to_numpy())
            for prefix, columns in items.items()
        },
        "kmo": float(kmo),
        "beta": beta,
    }


def assert_close(actual: float, expected: float, tolerance: float = 1e-10) -> None:
    if abs(actual - expected) > tolerance:
        raise AssertionError(f"{actual} != {expected}（容差 {tolerance}）")


def report_artifacts(directory: Path) -> list[Path]:
    names = {"analysis-result.json", "问卷数据分析报告.docx", "问卷数据分析报告.pdf"}
    return [path for path in directory.rglob("*") if path.is_file() and path.name in names]


def run_pipeline(
    *,
    name: str,
    source: Path,
    prompt: str,
    evidence_root: Path,
    plan_path: Path,
    seed: int,
    documents: bool,
    input_json: bool = False,
    expect_success: bool = True,
) -> tuple[dict[str, Any] | None, float, subprocess.CompletedProcess[str]]:
    output = evidence_root / name
    output.mkdir(parents=True, exist_ok=False)
    source_flag = "--input-json" if input_json else "--input"
    command = [
        sys.executable,
        str(PIPELINE),
        source_flag,
        str(source),
        "--plan",
        str(plan_path),
        "--output-dir",
        str(output),
        "--prompt",
        prompt,
        "--seed",
        str(seed),
    ]
    if not documents:
        command.append("--no-documents")
    environment = {
        **os.environ,
        "OPENBLAS_NUM_THREADS": "1",
        "OMP_NUM_THREADS": "1",
        "MKL_NUM_THREADS": "1",
    }
    started = time.perf_counter()
    completed = subprocess.run(
        command,
        cwd=ROOT,
        env=environment,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        timeout=240,
    )
    duration = time.perf_counter() - started
    log_dir = evidence_root / "logs"
    log_dir.mkdir(exist_ok=True)
    (log_dir / f"{name}.command.json").write_text(
        json.dumps({"command": command, "durationSeconds": duration, "returnCode": completed.returncode}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (log_dir / f"{name}.stdout.txt").write_text(completed.stdout, encoding="utf-8")
    (log_dir / f"{name}.stderr.txt").write_text(completed.stderr, encoding="utf-8")

    if expect_success and completed.returncode:
        raise AssertionError(completed.stderr or completed.stdout)
    if not expect_success:
        if completed.returncode == 0:
            raise AssertionError(f"{name} 应失败但返回成功")
        if report_artifacts(output):
            raise AssertionError(f"{name} 失败后生成了伪报告：{report_artifacts(output)}")
        return None, duration, completed
    result_path = output / "analysis-result.json"
    if not result_path.exists():
        raise AssertionError(f"{name} 未生成 analysis-result.json")
    return json.loads(result_path.read_text(encoding="utf-8")), duration, completed


def inspect_docx(docx_path: Path) -> dict[str, Any]:
    document = Document(docx_path)
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    section_titles = [
        paragraph.text for paragraph in document.paragraphs
        if re.match(r"^\d+、", paragraph.text)
    ]
    border_tables = 0
    for doc_table in document.tables:
        borders = doc_table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is not None and borders.find(qn("w:top")) is not None and borders.find(qn("w:bottom")) is not None:
            border_tables += 1
    with zipfile.ZipFile(docx_path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        document_xml = archive.read("word/document.xml").decode("utf-8")
    return {
        "paragraphText": paragraph_text,
        "sectionTitles": section_titles,
        "tableCount": len(document.tables),
        "threeLineTableCount": border_tables,
        "imageCount": len(media),
        "hasSongFont": "宋体" in document_xml,
        "hasHeiFont": "黑体" in document_xml,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="仅使用固定 synthetic 问卷重放交付主线")
    parser.add_argument("--sample", default=str(DEFAULT_SAMPLE))
    parser.add_argument("--csv", default=str(DEFAULT_CSV))
    parser.add_argument("--seed", type=int, default=20260712)
    parser.add_argument("--evidence-root")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    sample = assert_synthetic(Path(args.sample))
    csv_sample = assert_synthetic(Path(args.csv))
    run_stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    evidence_root = Path(args.evidence_root).resolve() if args.evidence_root else (
        ROOT / "data" / "validation" / "T-20260717-08-synthetic" / f"scientific-{run_stamp}"
    )
    if evidence_root.exists():
        raise SystemExit(f"证据目录已存在，拒绝覆盖：{evidence_root}")
    evidence_root.mkdir(parents=True)

    historical_before = hash_tree(OUTPUTS)
    (evidence_root / "historical-outputs-before.json").write_text(
        json.dumps(historical_before, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    plans_dir = evidence_root / "confirmed-plans"
    plans_dir.mkdir()
    full_plan = json.loads(PLAN.read_text(encoding="utf-8"))
    full_plan_path = plans_dir / "full.json"
    full_plan_path.write_text(json.dumps(full_plan, ensure_ascii=False, indent=2), encoding="utf-8")
    limited_plan = {
        **full_plan,
        "analysisOrder": ["data-quality", "reliability"],
        "confirmation": {"confirmed": True, "confirmedAtUtc": "2026-07-31T00:00:00.000Z"},
    }
    limited_plan_path = plans_dir / "reliability-only.json"
    limited_plan_path.write_text(json.dumps(limited_plan, ensure_ascii=False, indent=2), encoding="utf-8")
    invalid_plan = {
        **full_plan,
        "humanConfirmations": ["未识别到可执行的分析需求"],
        "confirmation": {"confirmed": False, "confirmedAtUtc": None},
    }
    invalid_plan_path = plans_dir / "unconfirmed-invalid-demand.json"
    invalid_plan_path.write_text(json.dumps(invalid_plan, ensure_ascii=False, indent=2), encoding="utf-8")

    full, full_seconds, _ = run_pipeline(
        name="full-xlsx",
        source=sample,
        prompt=FULL_PROMPT,
        evidence_root=evidence_root,
        plan_path=full_plan_path,
        seed=args.seed,
        documents=True,
    )
    repeat, repeat_seconds, _ = run_pipeline(
        name="full-repeat",
        source=sample,
        prompt=FULL_PROMPT,
        evidence_root=evidence_root,
        plan_path=full_plan_path,
        seed=args.seed,
        documents=False,
    )
    limited, limited_seconds, _ = run_pipeline(
        name="reliability-only",
        source=sample,
        prompt=LIMITED_PROMPT,
        evidence_root=evidence_root,
        plan_path=limited_plan_path,
        seed=args.seed,
        documents=True,
    )
    csv_result, csv_seconds, _ = run_pipeline(
        name="csv-reliability",
        source=csv_sample,
        prompt="仅做信度分析",
        evidence_root=evidence_root,
        plan_path=limited_plan_path,
        seed=args.seed,
        documents=False,
    )
    _, invalid_prompt_seconds, invalid_prompt_run = run_pipeline(
        name="invalid-demand",
        source=sample,
        prompt=INVALID_PROMPT,
        evidence_root=evidence_root,
        plan_path=invalid_plan_path,
        seed=args.seed,
        documents=True,
        expect_success=False,
    )

    frame = pd.read_excel(sample)
    missing_column_frame = frame.drop(columns=[next(column for column in frame.columns if str(column).startswith("AA3"))])
    fixtures_dir = evidence_root / "synthetic-invalid-fixtures"
    fixtures_dir.mkdir()
    missing_column_path = fixtures_dir / "synthetic-missing-AA3.json"
    missing_column_path.write_text(
        json.dumps({"headers": missing_column_frame.columns.tolist(), "rows": missing_column_frame.to_dict(orient="records")}, ensure_ascii=False),
        encoding="utf-8",
    )
    _, missing_column_seconds, missing_column_run = run_pipeline(
        name="missing-column",
        source=missing_column_path,
        prompt=FULL_PROMPT,
        evidence_root=evidence_root,
        plan_path=full_plan_path,
        seed=args.seed,
        documents=True,
        input_json=True,
        expect_success=False,
    )

    assert full is not None and repeat is not None and limited is not None and csv_result is not None
    reference = independent_reference(frame)
    assert_close(full["raw"]["kmo"], reference["kmo"])
    assert_close(full["raw"]["kmo"], SYNTHETIC_BENCHMARK["kmo"])
    for prefix, expected in reference["alphas"].items():
        assert_close(full["raw"]["reliability"][prefix], expected)
        assert_close(full["raw"]["reliability"][prefix], SYNTHETIC_BENCHMARK["alphas"][prefix])
    coefficient_names = ["const", "PP_c", "PC", "PI", "PV_c", "interaction"]
    for index, name in enumerate(coefficient_names):
        assert_close(full["raw"]["outcomeModel"][name], float(reference["beta"][index]))
        assert_close(full["raw"]["outcomeModel"][name], SYNTHETIC_BENCHMARK["outcomeModel"][name])
    if full["raw"]["indirect"] != repeat["raw"]["indirect"]:
        raise AssertionError("固定随机种子下 Bootstrap 间接效应不可复现")
    if full["raw"]["outcomeModel"] != repeat["raw"]["outcomeModel"]:
        raise AssertionError("固定输入下 OLS 结果不可复现")

    allowed_limited_titles = ["1、数据质量与变量识别", "2、信度检验"]
    actual_limited_titles = [
        payload["title"] for payload in limited["tables"]
        if payload["type"] == "text" and re.match(r"^\d+、", payload["title"])
    ]
    if actual_limited_titles != allowed_limited_titles:
        raise AssertionError(f"限定章节越界：{actual_limited_titles}")
    allowed_raw = {"planSummary", "analysisScope", "seed", "dataQuality", "effectiveSampleSizes", "reliability"}
    if set(limited["raw"]) != allowed_raw:
        raise AssertionError(f"限定 JSON 泄露未请求结果：{sorted(set(limited['raw']) - allowed_raw)}")
    limited_dir = evidence_root / "reliability-only"
    if (limited_dir / "研究模型图.png").exists() or (limited_dir / "调节效应图.png").exists():
        raise AssertionError("限定章节报告仍生成了模型图")
    if csv_result["raw"]["reliability"] != limited["raw"]["reliability"]:
        raise AssertionError("同一 synthetic 数据的 CSV/XLSX 信度结果不一致")

    full_docx = evidence_root / "full-xlsx" / "问卷数据分析报告.docx"
    full_pdf = evidence_root / "full-xlsx" / "问卷数据分析报告.pdf"
    limited_docx = limited_dir / "问卷数据分析报告.docx"
    limited_pdf = limited_dir / "问卷数据分析报告.pdf"
    for artifact in [full_docx, full_pdf, limited_docx, limited_pdf]:
        if not artifact.exists() or artifact.stat().st_size < 1000:
            raise AssertionError(f"报告产物缺失或异常：{artifact}")
    full_docx_check = inspect_docx(full_docx)
    limited_docx_check = inspect_docx(limited_docx)
    if full_docx_check["imageCount"] != 2:
        raise AssertionError(f"完整 DOCX 模型图数量异常：{full_docx_check['imageCount']}")
    if full_docx_check["tableCount"] != full_docx_check["threeLineTableCount"]:
        raise AssertionError("完整 DOCX 存在未通过三线表边框结构检查的表格")
    if not full_docx_check["hasSongFont"] or not full_docx_check["hasHeiFont"]:
        raise AssertionError("完整 DOCX 未检测到宋体/黑体中文字体声明")
    if limited_docx_check["sectionTitles"] != allowed_limited_titles:
        raise AssertionError(f"限定 DOCX 章节越界：{limited_docx_check['sectionTitles']}")
    if limited_docx_check["imageCount"] != 0:
        raise AssertionError("限定 DOCX 包含未请求图片")

    full_pdf_pages = len(pdfium.PdfDocument(full_pdf))
    limited_pdf_pages = len(pdfium.PdfDocument(limited_pdf))
    if not 5 <= full_pdf_pages <= 15:
        raise AssertionError(f"完整 PDF 页数异常：{full_pdf_pages}")
    if not 1 <= limited_pdf_pages < full_pdf_pages:
        raise AssertionError(f"限定 PDF 页数异常：{limited_pdf_pages}")

    historical_after = hash_tree(OUTPUTS)
    (evidence_root / "historical-outputs-after.json").write_text(
        json.dumps(historical_after, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if historical_before != historical_after:
        raise AssertionError("历史 outputs 前后哈希不一致")

    artifacts = hash_tree(evidence_root)
    summary = {
        "status": "passed",
        "syntheticOnly": True,
        "sample": str(sample),
        "csvSample": str(csv_sample),
        "seed": args.seed,
        "durationsSeconds": {
            "fullXlsx": full_seconds,
            "fullRepeat": repeat_seconds,
            "reliabilityOnly": limited_seconds,
            "csvReliability": csv_seconds,
            "invalidDemand": invalid_prompt_seconds,
            "missingColumn": missing_column_seconds,
        },
        "science": {
            "kmo": full["raw"]["kmo"],
            "alphas": full["raw"]["reliability"],
            "outcomeModel": full["raw"]["outcomeModel"],
            "bootstrapSamples": full["raw"]["bootstrapSamples"],
            "bootstrapReproducible": True,
        },
        "scope": {
            "request": LIMITED_PROMPT,
            "sectionTitles": actual_limited_titles,
            "rawKeys": sorted(limited["raw"]),
            "modelFigures": 0,
        },
        "errors": {
            "invalidDemandReturnCode": invalid_prompt_run.returncode,
            "invalidDemandMessage": invalid_prompt_run.stderr.strip().splitlines()[-1],
            "missingColumnReturnCode": missing_column_run.returncode,
            "missingColumnMessage": missing_column_run.stderr.strip().splitlines()[-1],
            "reportArtifactsCreated": 0,
        },
        "documents": {
            "fullDocx": {key: value for key, value in full_docx_check.items() if key != "paragraphText"},
            "limitedDocx": {key: value for key, value in limited_docx_check.items() if key != "paragraphText"},
            "fullPdfPages": full_pdf_pages,
            "limitedPdfPages": limited_pdf_pages,
        },
        "historicalOutputs": {
            "before": historical_before,
            "after": historical_after,
            "unchanged": historical_before == historical_after,
        },
        "artifactHashes": artifacts,
        "evidenceRoot": str(evidence_root),
    }
    summary_path = evidence_root / "machine-summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "status": "passed",
        "evidenceRoot": str(evidence_root),
        "machineSummary": str(summary_path),
        "kmo": summary["science"]["kmo"],
        "alphas": summary["science"]["alphas"],
        "fullPdfPages": full_pdf_pages,
        "limitedPdfPages": limited_pdf_pages,
        "historicalOutputsUnchanged": True,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
